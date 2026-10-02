#!/usr/bin/env python3
"""Build crosswalk/crosswalk.csv + crosswalk.xlsx from crosswalk/_src/r*.py and validate it.

Run from repo root:  python3 scripts/build_crosswalk.py
Exit code is non-zero if any validation fails.
"""
import csv, re, sys, importlib
from pathlib import Path
from collections import Counter

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "crosswalk" / "_src"))
sys.path.insert(0, str(ROOT / "scripts"))
from common import FIELDS  # noqa: E402
import refdata as R  # noqa: E402


def load_rows():
    rows = []
    for mod in ("r1", "r2", "r3", "r4", "r5"):
        rows += importlib.import_module(mod).ROWS
    return [dict(zip(FIELDS, r)) for r in rows]


def expand_ranges(text):
    """'5.24-5.27' -> 5.24..5.27 (27001 numbering only)."""
    out = set()
    for a, b in re.findall(r"\b([5-8]\.\d{1,2})-([5-8]\.\d{1,2})\b", text):
        pa, ia = a.split("."); pb, ib = b.split(".")
        if pa == pb:
            out |= {f"{pa}.{i}" for i in range(int(ia), int(ib) + 1)}
    return out


def validate(rows):
    errs = []
    ids = [r["obligation_id"] for r in rows]
    for i, c in Counter(ids).items():
        if c > 1:
            errs.append(f"duplicate obligation_id {i}")
    triples = Counter((r["eu_ai_act_article"], r["actor_role"], r["risk_tier"]) for r in rows)
    for t, c in triples.items():
        if c > 1:
            errs.append(f"duplicate (article, role, tier) triple {t}")
    for r in rows:
        oid = r["obligation_id"]
        for f in FIELDS:
            if not str(r[f]).strip():
                errs.append(f"{oid}: empty field {f}")
        if r["iso42001_verdict"] not in R.VERDICTS:
            errs.append(f"{oid}: bad iso42001_verdict {r['iso42001_verdict']!r}")
        if r["nist_verdict"] not in R.VERDICTS:
            errs.append(f"{oid}: bad nist_verdict {r['nist_verdict']!r}")
        if r["iso27001_verdict"] not in R.VERDICTS_27001:
            errs.append(f"{oid}: bad iso27001_verdict {r['iso27001_verdict']!r}")
        # ISO 42001 Annex A ids must exist in the reference list (which mirrors the supplied skeleton; see README provenance)
        for a in re.findall(r"\bA\.\d{1,2}(?:\.\d{1,2}){1,2}\b", r["iso42001_clause_or_control"]):
            if a not in R.ISO42001_ANNEX_A:
                errs.append(f"{oid}: unknown ISO 42001 Annex A id {a}")
        # No-coverage consistency
        if r["iso42001_verdict"] == "No coverage" and r["iso42001_clause_or_control"].strip() != "-":
            errs.append(f"{oid}: ISO verdict 'No coverage' but a reference is cited")
        if r["nist_verdict"] == "No coverage" and r["nist_ai_rmf_subcategory"].strip() != "-":
            errs.append(f"{oid}: NIST verdict 'No coverage' but a reference is cited")
        if r["iso42001_verdict"] != "No coverage" and r["iso42001_clause_or_control"].strip() == "-":
            errs.append(f"{oid}: ISO reference '-' but verdict is not 'No coverage'")
        if r["nist_verdict"] != "No coverage" and r["nist_ai_rmf_subcategory"].strip() == "-":
            errs.append(f"{oid}: NIST reference '-' but verdict is not 'No coverage'")
        # NIST ids exist
        if r["nist_ai_rmf_subcategory"].strip() != "-":
            for tok in [t.strip() for t in r["nist_ai_rmf_subcategory"].split(";")]:
                if tok not in R.NIST_AI_RMF:
                    errs.append(f"{oid}: unknown NIST AI RMF subcategory {tok!r}")
        # Main-body clause numbers are NOT cited anywhere (they could not be verified against a licensed copy).
        for col in ("iso42001_clause_or_control", "iso27001_2022_supplementary"):
            if re.search(r"\b(?:cl\.|clauses?)\s*\d", r[col], flags=re.I):
                errs.append(f"{oid}: main-body clause number cited in {col}; cite Annex A controls only")
        # 27001 ids exist
        txt = r["iso27001_2022_supplementary"]
        found = set(re.findall(r"\b([5-8]\.\d{1,2})\b", txt)) | expand_ranges(txt)
        for a in found:
            if a not in R.ISO27001_ANNEX_A:
                errs.append(f"{oid}: unknown ISO 27001 Annex A id {a}")
        if r["iso27001_verdict"] == "Not assessed" and r["iso27001_2022_supplementary"].strip() != "-":
            errs.append(f"{oid}: 27001 'Not assessed' but a reference is cited")
        # Copyright guard (heuristic): author-written intent only, short, no quotation marks
        cell = r["iso42001_clause_or_control"]
        if len(cell) > 420:
            errs.append(f"{oid}: ISO 42001 cell > 420 chars (possible pasted text)")
        if '"' in cell or "“" in cell:
            errs.append(f"{oid}: quotation marks in ISO 42001 cell (possible verbatim text)")
    return errs


def write_csv(rows, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)


def write_xlsx(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Crosswalk"
    fills = {"Substantially evidences": "C6E0B4", "Partially evidences": "FFE699",
             "Adjacent": "F8CBAD", "No coverage": "E06666", "Not assessed": "D9D9D9"}
    thin = Side(style="thin", color="BFBFBF")
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    ws.append(FIELDS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append([r[f] for f in FIELDS])
    widths = [20, 18, 55, 14, 14, 30, 34, 55, 16, 40, 14, 26, 14, 60, 45, 38]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    vcols = [FIELDS.index(k) + 1 for k in ("iso42001_verdict", "iso27001_verdict", "nist_verdict")]
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if c.column in vcols and c.value in fills:
                c.fill = PatternFill("solid", fgColor=fills[c.value])
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    # Summary sheet
    s = wb.create_sheet("Summary")
    s.append(["Verdict distribution (ISO/IEC 42001 column)"]); s["A1"].font = Font(bold=True)
    c42 = Counter(r["iso42001_verdict"] for r in rows)
    cn = Counter(r["nist_verdict"] for r in rows)
    s.append(["Verdict", "ISO 42001 rows", "NIST AI RMF rows"])
    for v in ("Substantially evidences", "Partially evidences", "Adjacent", "No coverage"):
        s.append([v, c42.get(v, 0), cn.get(v, 0)])
    s.append([]); s.append(["Rows by actor role"]); s.cell(s.max_row, 1).font = Font(bold=True)
    for k, v in Counter(r["actor_role"] for r in rows).items():
        s.append([k, v])
    s.append([]); s.append(["Rows by risk tier"]); s.cell(s.max_row, 1).font = Font(bold=True)
    for k, v in Counter(r["risk_tier"].split(" (")[0] for r in rows).items():
        s.append([k, v])
    s.column_dimensions["A"].width = 34; s.column_dimensions["B"].width = 18; s.column_dimensions["C"].width = 18

    # Exposure sheet for an ISO 42001-certified organisation
    g = wb.create_sheet("ISO-certified gap view")
    g.append(["obligation_id", "eu_ai_act_article", "actor_role", "risk_tier", "iso42001_verdict",
              "exposure_reading", "residual_gap"])
    for c in g[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr_fill
    reading = {"Substantially evidences": "Closest evidence - still verify the named gap",
               "Partially evidences": "Evidence helps - named gap remains",
               "Adjacent": "Exposed - certificate does not evidence this",
               "No coverage": "Exposed - nothing in ISO 42001 speaks to this"}
    order = {"No coverage": 0, "Adjacent": 1, "Partially evidences": 2, "Substantially evidences": 3}
    for r in sorted(rows, key=lambda x: (order[x["iso42001_verdict"]], x["obligation_id"])):
        g.append([r["obligation_id"], r["eu_ai_act_article"], r["actor_role"], r["risk_tier"].split(" (")[0],
                  r["iso42001_verdict"], reading[r["iso42001_verdict"]], r["residual_gap"]])
    for i, w in enumerate([22, 24, 18, 16, 22, 44, 90], 1):
        g.column_dimensions[get_column_letter(i)].width = w
    for row in g.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[4].fill = PatternFill("solid", fgColor=fills[row[4].value])
    g.freeze_panes = "A2"
    wb.save(path)


def main():
    rows = load_rows()
    errs = validate(rows)
    if errs:
        print("VALIDATION FAILED:")
        for e in errs:
            print(" -", e)
        sys.exit(1)
    out = ROOT / "crosswalk"
    write_csv(rows, out / "crosswalk.csv")
    write_xlsx(rows, out / "crosswalk.xlsx")
    c = Counter(r["iso42001_verdict"] for r in rows)
    print(f"OK: {len(rows)} rows; ISO 42001 verdicts: {dict(c)}")


if __name__ == "__main__":
    main()
