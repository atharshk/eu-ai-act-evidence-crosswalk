#!/usr/bin/env python3
"""Build register/risk_register.xlsx (live formulas) and register/risk_register_worked_examples.csv.

Run AFTER build_crosswalk.py:  python3 scripts/build_register.py
The register links to crosswalk/crosswalk.csv through the 'Crosswalk IDs' sheet; Actor Role, Risk Tier and the
ISO verdict are looked up by formula from the Linked Obligation ID so they cannot drift from the crosswalk.
"""
import csv, datetime as dt, sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import scoring as S  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

HEAD = [
    "Risk ID", "Linked Crosswalk Obligation ID", "AI System / Use Case", "Actor Role (auto)", "Risk Tier (auto)",
    "Risk Description", "Affected Party (individuals / org / both)",
    "Likelihood (1-5)", "Rights Impact (1-5)", "Organisational Impact (1-5)",
    "Rights Impact (L x I)", "Organisational Impact (L x I)", "Inherent Score", "Score driver",
    "Existing Controls (ISO 42001 ref)", "Crosswalk ISO 42001 verdict (auto)",
    "Control Effectiveness (as rated)", "Effectiveness cap from verdict (auto)",
    "Override justification (non-ISO control)", "Effectiveness applied",
    "Residual Likelihood", "Residual Rights (L x I)", "Residual Org (L x I)", "Residual Score", "Residual band",
    "Risk Owner (role)", "Treatment Decision", "Target Date", "Evidence Reference", "Review Date",
]
COL = {h: get_column_letter(i + 1) for i, h in enumerate(HEAD)}

EXAMPLES = [
    dict(id="RISK-001", ob="PRV-HR-A10-BIAS",
         sys="EXAMPLE (fictional): recruitment CV-screening and candidate-ranking tool sold to employers (Annex III 4(a)).",
         desc="Training and validation data under-represent candidates with career breaks and some demographic groups, so ranking systematically scores them lower. Non-conformity with Art. 10(2)(f)-(g) and 10(3)-(4); outcome is discrimination in access to employment.",
         party="both", L=3, RI=4, OI=3,
         ctrl="A.7.4 Quality of data; A.5.4 Impact on individuals or groups; A.6.2.4 Verification and validation.",
         eff="Partially effective", ovr="",
         owner="Head of AI Product (accountable); Chief Data Officer (responsible); DPO (consulted)",
         treat="Mitigate", target=dt.date(2027, 6, 30),
         evid="EV-A10-BIAS-001 (placeholder): bias evaluation report v1.2 - run once on gender only; no test on employment gaps; no re-test cadence.",
         review=dt.date(2027, 1, 15)),
    dict(id="RISK-002", ob="GPP-GP-A53-COPY",
         sys="EXAMPLE (fictional): general-purpose language model offered to downstream providers through an API.",
         desc="Data-acquisition pipeline does not reliably detect or honour machine-readable rights reservations (Art. 4(3), Directive (EU) 2019/790), so reserved works enter training data and the Art. 53(1)(c) copyright policy cannot be evidenced to the AI Office.",
         party="both", L=3, RI=2, OI=4,
         ctrl="A.7.3 Acquisition of data; A.7.5 Data provenance (these give provenance records but do not address rights reservations).",
         eff="Partially effective",
         ovr="Non-ISO control: opt-out detection filter in the crawler with a monthly sampled test log. Bespoke technical control, not an ISO 42001 control; test log is the evidence.",
         owner="General Counsel (accountable); Head of Data Acquisition (responsible)",
         treat="Mitigate", target=dt.date(2026, 12, 31),
         evid="EV-A53-COPY-001 (placeholder): crawler opt-out filter test log.",
         review=dt.date(2026, 11, 15)),
]


def load_crosswalk():
    with open(ROOT / "crosswalk" / "crosswalk.csv", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def py_calc(L, RI, OI, verdict, eff, ovr):
    ranks = [e[0] for e in S.EFFECTIVENESS]
    cap = S.VERDICT_CAP[verdict]
    applied = cap if (ranks.index(eff) > ranks.index(cap) and not ovr) else eff
    resL = max(1, L - ranks.index(applied))
    rr, ro = resL * RI, resL * OI
    x = max(rr, ro)
    base = 4 if x >= 16 else 3 if x >= 10 else 2 if x >= 5 else 1
    floor = 3 if (RI == 5 or (RI >= 4 and resL >= 2)) else 1
    band = ["Low", "Medium", "High", "Critical"][max(base, floor) - 1]
    return dict(rights=L * RI, org=L * OI, inherent=max(L * RI, L * OI), cap=cap, applied=applied,
                resL=resL, res_rights=rr, res_org=ro, residual=x, band=band)


def main():
    cw = load_crosswalk()
    byid = {r["obligation_id"]: r for r in cw}
    wb = Workbook()
    thin = Side(style="thin", color="BFBFBF")
    hdr = PatternFill("solid", fgColor="1F3864")
    auto = PatternFill("solid", fgColor="EDEDED")

    # ---------- Read me ----------
    rm = wb.active; rm.title = "Read me"
    lines = [
        ("Risk register layer - EU AI Act evidence crosswalk", True),
        ("Every risk row must cite a Linked Crosswalk Obligation ID. Actor role, risk tier and the ISO 42001 verdict are looked up from the crosswalk by formula.", False),
        ("'Risk' is scored twice: Rights Impact (harm to health, safety and fundamental rights of persons - the Act's sense) and Organisational Impact (harm to the organisation - the enterprise-risk sense). They do not always point the same way.", False),
        ("Grey columns are formulas - do not overwrite. Scales, effectiveness caps, bands and limitations are on the 'Scoring Methodology' sheet.", False),
        ("Rows RISK-001 and RISK-002 are ILLUSTRATIVE worked examples about fictional systems. Evidence references are placeholders, not real documents.", False),
        ("Dates were set relative to 2 Oct 2026 and to the application dates in the crosswalk (as of 2 Oct 2026; subject to further legislative amendment).", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        rm.cell(i, 1, t).font = Font(bold=b, size=13 if b else 11)
        rm.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
    rm.column_dimensions["A"].width = 130

    # ---------- Lists ----------
    ls = wb.create_sheet("Lists")
    ls.append(["Effectiveness", "Likelihood reduction", "Definition", "", "Crosswalk verdict", "Max effectiveness"])
    for i, (n, red, d) in enumerate(S.EFFECTIVENESS):
        ls.cell(i + 2, 1, n); ls.cell(i + 2, 2, red); ls.cell(i + 2, 3, d)
    for i, (v, cap) in enumerate(S.VERDICT_CAP.items()):
        ls.cell(i + 2, 5, v); ls.cell(i + 2, 6, cap)
    ls["H1"] = "Treatment"; [ls.cell(i + 2, 8, t) for i, t in enumerate(["Mitigate", "Accept", "Transfer", "Avoid"])]
    ls["J1"] = "Affected party"; [ls.cell(i + 2, 10, t) for i, t in enumerate(["individuals", "org", "both"])]
    for c in ls[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCEFHJ", [22, 20, 70, 24, 22, 14, 16]):
        ls.column_dimensions[col].width = w

    # ---------- Crosswalk IDs ----------
    ci = wb.create_sheet("Crosswalk IDs")
    ci.append(["obligation_id", "eu_ai_act_article", "actor_role", "risk_tier", "iso42001_verdict", "nist_verdict"])
    for c in ci[1]:
        c.font = Font(bold=True)
    for r in cw:
        ci.append([r["obligation_id"], r["eu_ai_act_article"], r["actor_role"], r["risk_tier"].split(" (")[0],
                   r["iso42001_verdict"], r["nist_verdict"]])
    for i, w in enumerate([22, 26, 20, 22, 24, 24], 1):
        ci.column_dimensions[get_column_letter(i)].width = w
    n_cw = len(cw) + 1

    # ---------- Register ----------
    ws = wb.create_sheet("Register", 1)
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 48
    LAST = 30
    c = COL

    def rowformulas(r):
        return {
            "Actor Role (auto)": f'=IF({c["Linked Crosswalk Obligation ID"]}{r}="","",IFERROR(INDEX(\'Crosswalk IDs\'!$C$2:$C${n_cw},MATCH({c["Linked Crosswalk Obligation ID"]}{r},\'Crosswalk IDs\'!$A$2:$A${n_cw},0)),"ID NOT IN CROSSWALK"))',
            "Risk Tier (auto)": f'=IF({c["Linked Crosswalk Obligation ID"]}{r}="","",IFERROR(INDEX(\'Crosswalk IDs\'!$D$2:$D${n_cw},MATCH({c["Linked Crosswalk Obligation ID"]}{r},\'Crosswalk IDs\'!$A$2:$A${n_cw},0)),"ID NOT IN CROSSWALK"))',
            "Rights Impact (L x I)": f'=IF(OR({c["Likelihood (1-5)"]}{r}="",{c["Rights Impact (1-5)"]}{r}=""),"",{c["Likelihood (1-5)"]}{r}*{c["Rights Impact (1-5)"]}{r})',
            "Organisational Impact (L x I)": f'=IF(OR({c["Likelihood (1-5)"]}{r}="",{c["Organisational Impact (1-5)"]}{r}=""),"",{c["Likelihood (1-5)"]}{r}*{c["Organisational Impact (1-5)"]}{r})',
            "Inherent Score": f'=IF({c["Rights Impact (L x I)"]}{r}="","",MAX({c["Rights Impact (L x I)"]}{r},{c["Organisational Impact (L x I)"]}{r}))',
            "Score driver": f'=IF({c["Inherent Score"]}{r}="","",IF({c["Rights Impact (L x I)"]}{r}>{c["Organisational Impact (L x I)"]}{r},"Rights",IF({c["Organisational Impact (L x I)"]}{r}>{c["Rights Impact (L x I)"]}{r},"Organisation","Both equal")))',
            "Crosswalk ISO 42001 verdict (auto)": f'=IF({c["Linked Crosswalk Obligation ID"]}{r}="","",IFERROR(INDEX(\'Crosswalk IDs\'!$E$2:$E${n_cw},MATCH({c["Linked Crosswalk Obligation ID"]}{r},\'Crosswalk IDs\'!$A$2:$A${n_cw},0)),"ID NOT IN CROSSWALK"))',
            "Effectiveness cap from verdict (auto)": f'=IF({c["Crosswalk ISO 42001 verdict (auto)"]}{r}="","",IFERROR(INDEX(Lists!$F$2:$F$5,MATCH({c["Crosswalk ISO 42001 verdict (auto)"]}{r},Lists!$E$2:$E$5,0)),""))',
            "Effectiveness applied": f'=IF(OR({c["Control Effectiveness (as rated)"]}{r}="",{c["Effectiveness cap from verdict (auto)"]}{r}=""),"",IF(AND(MATCH({c["Control Effectiveness (as rated)"]}{r},Lists!$A$2:$A$4,0)>MATCH({c["Effectiveness cap from verdict (auto)"]}{r},Lists!$A$2:$A$4,0),{c["Override justification (non-ISO control)"]}{r}=""),{c["Effectiveness cap from verdict (auto)"]}{r},{c["Control Effectiveness (as rated)"]}{r}))',
            "Residual Likelihood": f'=IF({c["Effectiveness applied"]}{r}="","",MAX(1,{c["Likelihood (1-5)"]}{r}-(MATCH({c["Effectiveness applied"]}{r},Lists!$A$2:$A$4,0)-1)))',
            "Residual Rights (L x I)": f'=IF({c["Residual Likelihood"]}{r}="","",{c["Residual Likelihood"]}{r}*{c["Rights Impact (1-5)"]}{r})',
            "Residual Org (L x I)": f'=IF({c["Residual Likelihood"]}{r}="","",{c["Residual Likelihood"]}{r}*{c["Organisational Impact (1-5)"]}{r})',
            "Residual Score": f'=IF({c["Residual Likelihood"]}{r}="","",MAX({c["Residual Rights (L x I)"]}{r},{c["Residual Org (L x I)"]}{r}))',
            "Residual band": f'=IF({c["Residual Score"]}{r}="","",CHOOSE(MAX(IF({c["Residual Score"]}{r}>=16,4,IF({c["Residual Score"]}{r}>=10,3,IF({c["Residual Score"]}{r}>=5,2,1))),IF(OR({c["Rights Impact (1-5)"]}{r}=5,AND({c["Rights Impact (1-5)"]}{r}>=4,{c["Residual Likelihood"]}{r}>=2)),3,1)),"Low","Medium","High","Critical"))',
        }

    ex_by_row = {2 + i: e for i, e in enumerate(EXAMPLES)}
    for r in range(2, LAST + 1):
        f = rowformulas(r)
        e = ex_by_row.get(r)
        vals = {}
        if e:
            vals = {"Risk ID": e["id"], "Linked Crosswalk Obligation ID": e["ob"], "AI System / Use Case": e["sys"],
                    "Risk Description": e["desc"], "Affected Party (individuals / org / both)": e["party"],
                    "Likelihood (1-5)": e["L"], "Rights Impact (1-5)": e["RI"], "Organisational Impact (1-5)": e["OI"],
                    "Existing Controls (ISO 42001 ref)": e["ctrl"], "Control Effectiveness (as rated)": e["eff"],
                    "Override justification (non-ISO control)": e["ovr"] or None, "Risk Owner (role)": e["owner"],
                    "Treatment Decision": e["treat"], "Target Date": e["target"], "Evidence Reference": e["evid"],
                    "Review Date": e["review"]}
        for h in HEAD:
            cell = ws[f"{COL[h]}{r}"]
            if h in f:
                cell.value = f[h]; cell.fill = auto
            elif h in vals:
                cell.value = vals[h]
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if h in ("Target Date", "Review Date"):
                cell.number_format = "yyyy-mm-dd"

    dv_n = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                          showErrorMessage=True, error="Enter a whole number 1-5 (see Scoring Methodology).")
    dv_e = DataValidation(type="list", formula1="=Lists!$A$2:$A$4", allow_blank=True)
    dv_t = DataValidation(type="list", formula1="=Lists!$H$2:$H$5", allow_blank=True)
    dv_p = DataValidation(type="list", formula1="=Lists!$J$2:$J$4", allow_blank=True)
    dv_o = DataValidation(type="list", formula1=f"='Crosswalk IDs'!$A$2:$A${n_cw}", allow_blank=True,
                          showErrorMessage=True, error="Must be an obligation_id from the crosswalk.")
    for dv in (dv_n, dv_e, dv_t, dv_p, dv_o):
        ws.add_data_validation(dv)
    for h in ("Likelihood (1-5)", "Rights Impact (1-5)", "Organisational Impact (1-5)"):
        dv_n.add(f"{COL[h]}2:{COL[h]}{LAST}")
    dv_e.add(f"{COL['Control Effectiveness (as rated)']}2:{COL['Control Effectiveness (as rated)']}{LAST}")
    dv_t.add(f"{COL['Treatment Decision']}2:{COL['Treatment Decision']}{LAST}")
    dv_p.add(f"{COL['Affected Party (individuals / org / both)']}2:{COL['Affected Party (individuals / org / both)']}{LAST}")
    dv_o.add(f"{COL['Linked Crosswalk Obligation ID']}2:{COL['Linked Crosswalk Obligation ID']}{LAST}")
    widths = [11, 22, 38, 14, 14, 55, 14, 10, 10, 12, 11, 12, 10, 13, 40, 20, 18, 20, 40, 18, 10, 11, 11, 10, 11, 36, 12, 12, 40, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"

    # ---------- Scoring Methodology ----------
    sm = wb.create_sheet("Scoring Methodology", 2)
    sm.column_dimensions["A"].width = 10; sm.column_dimensions["B"].width = 20; sm.column_dimensions["C"].width = 120
    r = 1

    def title(t):
        nonlocal r
        sm.cell(r, 1, t).font = Font(bold=True, size=12); r += 1

    def table(head, rows_):
        nonlocal r
        for j, h in enumerate(head, 1):
            cc = sm.cell(r, j, h); cc.font = Font(bold=True, color="FFFFFF"); cc.fill = hdr
        r += 1
        for row in rows_:
            for j, v in enumerate(row, 1):
                sm.cell(r, j, v).alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
        r += 1

    title("1. Why two impact scales")
    sm.cell(r, 1, "In enterprise risk management 'risk' means risk to the organisation. In the EU AI Act 'risk' means risk to the health, safety and fundamental rights of persons. They are different objects and can diverge: a system can be commercially low-risk and high-risk to the people it affects. The register therefore scores Rights Impact and Organisational Impact separately and shows which one drives the result.").alignment = Alignment(wrap_text=True, vertical="top")
    sm.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3); sm.row_dimensions[r].height = 62; r += 2
    title("2. Likelihood (next 12 months, in the system's operating context)"); table(["Score", "Label", "Definition"], S.LIKELIHOOD)
    title("3. Rights Impact (harm to persons)"); table(["Score", "Label", "Definition"], S.RIGHTS_IMPACT)
    title("4. Organisational Impact (harm to the organisation)"); table(["Score", "Label", "Definition"], S.ORG_IMPACT)
    title("5. Control effectiveness (rated on evidence of operation)"); table(["Reduction", "Rating", "Definition"], [(red, n, d) for n, red, d in S.EFFECTIVENESS])
    title("6. Effectiveness cap from the crosswalk verdict")
    table(["", "Crosswalk ISO 42001 verdict", "Maximum effectiveness unless a written override cites a non-ISO control"],
          [("", v, cap) for v, cap in S.VERDICT_CAP.items()])
    title("7. Formulae")
    table(["", "Item", "Rule"], [
        ("", "Rights score / Org score", "Likelihood x Rights Impact ; Likelihood x Organisational Impact"),
        ("", "Inherent score", "MAX(Rights score, Organisational score); driver column shows which one"),
        ("", "Residual likelihood", "MAX(1, Likelihood - reduction for the applied effectiveness)"),
        ("", "Residual score", "MAX(Residual likelihood x Rights Impact, Residual likelihood x Org Impact). Impacts are not reduced by controls in this version."),
        ("", "Bands (by product)", "; ".join(f"{n} {a}-{b}" for n, a, b in S.BANDS)),
        ("", "Rights floor", S.RIGHTS_OVERRIDE)])
    title("8. Limitations (named by the author)")
    table(["#", "", "Limitation"], [(i + 1, "", t) for i, t in enumerate(S.LIMITATIONS)])
    wb.save(ROOT / "register" / "risk_register.xlsx")

    # ---------- CSV of worked examples (values computed in Python; cross-checked vs LibreOffice) ----------
    out = ROOT / "register" / "risk_register_worked_examples.csv"
    with open(out, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(HEAD)
        for e in EXAMPLES:
            cwr = byid[e["ob"]]
            k = py_calc(e["L"], e["RI"], e["OI"], cwr["iso42001_verdict"], e["eff"], e["ovr"])
            w.writerow([e["id"], e["ob"], e["sys"], cwr["actor_role"], cwr["risk_tier"].split(" (")[0], e["desc"], e["party"],
                        e["L"], e["RI"], e["OI"], k["rights"], k["org"], k["inherent"],
                        "Rights" if k["rights"] > k["org"] else "Organisation" if k["org"] > k["rights"] else "Both equal",
                        e["ctrl"], cwr["iso42001_verdict"], e["eff"], k["cap"], e["ovr"], k["applied"],
                        k["resL"], k["res_rights"], k["res_org"], k["residual"], k["band"],
                        e["owner"], e["treat"], e["target"].isoformat(), e["evid"], e["review"].isoformat()])
            print(e["id"], k)


if __name__ == "__main__":
    main()
